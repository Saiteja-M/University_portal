from collections import defaultdict
from io import BytesIO

from django.db.models import Count, Q
from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.academics.models import AcademicYear, Course, Program, Semester
from apps.students.models import Enrollment, Student

from .models import AttendanceRecord, AttendanceSession


def _attendance_percentage(
    present: int,
    late: int,
    total: int,
) -> float:
    if total == 0:
        return 0.0

    return round(
        ((present + late) / total) * 100,
        2,
    )


def get_class_attendance_report(
    *,
    program_id: int,
    academic_year_id: int,
    semester_id: int,
    course_id: int,
):
    """
    Build the complete attendance dataset for one class/course.

    Class identity:
        Program
        Academic Year
        Semester
        Course

    Attendance percentages are calculated from AttendanceRecord.
    """

    program = Program.objects.select_related(
        "department",
    ).get(
        pk=program_id,
    )

    academic_year = AcademicYear.objects.get(
        pk=academic_year_id,
    )

    semester = Semester.objects.select_related(
        "program",
        "academic_year",
    ).get(
        pk=semester_id,
    )

    course = Course.objects.select_related(
        "semester",
        "semester__program",
        "regulation",
    ).get(
        pk=course_id,
    )

    # ---------------------------------------------------------
    # Validate academic relationships
    # ---------------------------------------------------------

    if semester.program_id != program.id:
        raise ValueError(
            "The selected semester does not belong "
            "to the selected program."
        )

    if semester.academic_year_id != academic_year.id:
        raise ValueError(
            "The selected semester does not belong "
            "to the selected academic year."
        )

    if course.semester_id != semester.id:
        raise ValueError(
            "The selected course does not belong "
            "to the selected semester."
        )

    # ---------------------------------------------------------
    # Determine year of study
    # ---------------------------------------------------------

    year_of_study = (
        (semester.number + 1) // 2
    )

    # ---------------------------------------------------------
    # Students in the selected class
    # ---------------------------------------------------------

    students = list(
        Student.objects
        .select_related(
            "user",
            "program",
        )
        .filter(
            program_id=program.id,
            status=Student.Status.ACTIVE,
            enrollments__academic_year_id=academic_year.id,
            enrollments__semester_id=semester.id,
            enrollments__status=Enrollment.Status.ACTIVE,
        )
        .distinct()
        .order_by(
            "student_id",
        )
    )

    # ---------------------------------------------------------
    # Sessions for selected course/class
    # ---------------------------------------------------------

    sessions = list(
        AttendanceSession.objects
        .filter(
            course_id=course.id,
            academic_year_id=academic_year.id,
            semester_id=semester.id,
        )
        .select_related(
            "faculty",
            "faculty__user",
        )
        .order_by(
            "session_date",
            "period",
        )
    )

    session_ids = [
        session.id
        for session in sessions
    ]

    # ---------------------------------------------------------
    # Attendance records
    # ---------------------------------------------------------

    records = (
        AttendanceRecord.objects
        .filter(
            session_id__in=session_ids,
            student_id__in=[
                student.id
                for student in students
            ],
        )
        .select_related(
            "student",
            "session",
        )
    )

    record_map = {}

    for record in records:
        record_map[
            (
                record.student_id,
                record.session_id,
            )
        ] = record

    # ---------------------------------------------------------
    # Build student summaries
    # ---------------------------------------------------------

    student_rows = []

    for student in students:
        present = 0
        absent = 0
        late = 0

        daily = {}

        for session in sessions:
            record = record_map.get(
                (
                    student.id,
                    session.id,
                )
            )

            status = (
                record.status
                if record
                else None
            )

            if status == "PRESENT":
                present += 1

            elif status == "ABSENT":
                absent += 1

            elif status == "LATE":
                late += 1

            daily[session.id] = status

        total = len(sessions)

        percentage = _attendance_percentage(
            present,
            late,
            total,
        )

        student_rows.append(
            {
                "student_id": student.student_id,
                "student_name": (
                    student.user.get_full_name()
                    or student.user.username
                ),
                "total_classes": total,
                "present": present,
                "absent": absent,
                "late": late,
                "attendance_percentage": percentage,
                "daily": daily,
            }
        )

    return {
        "program": program,
        "academic_year": academic_year,
        "semester": semester,
        "course": course,
        "year_of_study": year_of_study,
        "sessions": sessions,
        "students": student_rows,
    }


def build_excel_report(report):
    """
    Generate a professional multi-sheet Excel workbook.
    """

    workbook = Workbook()

    # =========================================================
    # Sheet 1: Class Summary
    # =========================================================

    summary = workbook.active
    summary.title = "Class Summary"

    summary["A1"] = "UNIVERSITY PORTAL"
    summary["A1"].font = Font(
        bold=True,
        size=18,
    )

    summary["A2"] = "Attendance Report"
    summary["A2"].font = Font(
        bold=True,
        size=14,
    )

    metadata = [
        (
            "Branch / Department",
            report["program"].department.name,
        ),
        (
            "Program",
            report["program"].name,
        ),
        (
            "Academic Year",
            report["academic_year"].name,
        ),
        (
            "Year of Study",
            f'Year {report["year_of_study"]}',
        ),
        (
            "Semester",
            f'Semester {report["semester"].number}',
        ),
        (
            "Course",
            (
                f'{report["course"].code} - '
                f'{report["course"].name}'
            ),
        ),
    ]

    row = 4

    for label, value in metadata:
        summary.cell(
            row=row,
            column=1,
            value=label,
        ).font = Font(
            bold=True,
        )

        summary.cell(
            row=row,
            column=2,
            value=value,
        )

        row += 1

    row += 1

    headers = [
        "Student ID",
        "Student Name",
        "Total Classes",
        "Present",
        "Absent",
        "Late",
        "Attendance %",
    ]

    for column, header in enumerate(
        headers,
        start=1,
    ):
        cell = summary.cell(
            row=row,
            column=column,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    row += 1

    for student in report["students"]:
        values = [
            student["student_id"],
            student["student_name"],
            student["total_classes"],
            student["present"],
            student["absent"],
            student["late"],
            student["attendance_percentage"] / 100,
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            cell = summary.cell(
                row=row,
                column=column,
                value=value,
            )

            if column == 7:
                cell.number_format = "0.00%"

        row += 1

    # =========================================================
    # Sheet 2: Daily Register
    # =========================================================

    daily = workbook.create_sheet(
        "Daily Register",
    )

    daily_headers = [
        "Student ID",
        "Student Name",
    ]

    for session in report["sessions"]:
        daily_headers.append(
            (
                session.session_date.strftime(
                    "%d-%b"
                )
                + f" P{session.period}"
            )
        )

    daily_headers.extend(
        [
            "Total",
            "Present",
            "Absent",
            "Late",
            "Attendance %",
        ]
    )

    for column, header in enumerate(
        daily_headers,
        start=1,
    ):
        cell = daily.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    row = 2

    for student in report["students"]:
        daily.cell(
            row=row,
            column=1,
            value=student["student_id"],
        )

        daily.cell(
            row=row,
            column=2,
            value=student["student_name"],
        )

        column = 3

        for session in report["sessions"]:
            status = student["daily"].get(
                session.id
            )

            display_status = {
                "PRESENT": "P",
                "ABSENT": "A",
                "LATE": "L",
            }.get(
                status,
                "-",
            )

            daily.cell(
                row=row,
                column=column,
                value=display_status,
            )

            column += 1

        daily.cell(
            row=row,
            column=column,
            value=student["total_classes"],
        )

        daily.cell(
            row=row,
            column=column + 1,
            value=student["present"],
        )

        daily.cell(
            row=row,
            column=column + 2,
            value=student["absent"],
        )

        daily.cell(
            row=row,
            column=column + 3,
            value=student["late"],
        )

        percentage_cell = daily.cell(
            row=row,
            column=column + 4,
            value=(
                student[
                    "attendance_percentage"
                ] / 100
            ),
        )

        percentage_cell.number_format = (
            "0.00%"
        )

        row += 1

    # =========================================================
    # Sheet 3: Sessions
    # =========================================================

    session_sheet = workbook.create_sheet(
        "Sessions",
    )

    session_headers = [
        "Date",
        "Period",
        "Course Code",
        "Course Name",
        "Topic",
        "Faculty",
    ]

    for column, header in enumerate(
        session_headers,
        start=1,
    ):
        cell = session_sheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

    for row, session in enumerate(
        report["sessions"],
        start=2,
    ):
        session_sheet.cell(
            row=row,
            column=1,
            value=session.session_date,
        )

        session_sheet.cell(
            row=row,
            column=2,
            value=session.period,
        )

        session_sheet.cell(
            row=row,
            column=3,
            value=report["course"].code,
        )

        session_sheet.cell(
            row=row,
            column=4,
            value=report["course"].name,
        )

        session_sheet.cell(
            row=row,
            column=5,
            value=session.topic,
        )

        session_sheet.cell(
            row=row,
            column=6,
            value=(
                session.faculty.user.get_full_name()
                or session.faculty.user.username
            ),
        )

    # =========================================================
    # Formatting
    # =========================================================

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"

        for column_cells in worksheet.columns:
            max_length = 0

            for cell in column_cells:
                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            column_letter = get_column_letter(
                column_cells[0].column
            )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                40,
            )

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


def download_excel_response(
    report,
):
    workbook = build_excel_report(
        report,
    )

    filename = (
        f'{report["program"].code}_'
        f'Sem{report["semester"].number}_'
        f'{report["course"].code}_'
        f'Attendance.xlsx'
    )

    response = HttpResponse(
        workbook.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response